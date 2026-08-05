"""
Tests for the sandboxed out-of-process document conversion service.

MarkItDown used to parse untrusted documents in-process (chat attachments in
overlord.py and knowledge sources in agents/knowledge/base.py): a hostile file
could hang, crash, or balloon the runtime. document_converter now runs every
conversion in a spawned subprocess with rlimits and a wall-clock kill, routes
PDFs to pdf-inspector, and returns typed QuarantineReason outcomes instead of
raising.

These tests use real converters and real hostile fixtures - no mocks:
  - a decompression-bomb .docx (tiny zip, enormous extracted text)
  - a huge HTML document that cannot finish inside the timeout
  - oversize input rejected before any subprocess is spawned
  - PDFs routed to pdf-inspector, with fallback to MarkItDown when
    pdf-inspector cannot parse the file
  - office/OpenDocument/RTF/EPUB/CSV routed to anydoc, with fallback to
    MarkItDown only for the formats MarkItDown supports; anydoc-only formats
    quarantine honestly as parser_error
"""

import asyncio
import io
import threading
import zipfile

import pytest
from fpdf import FPDF
from markitdown import MarkItDown
from pypdf import PdfReader, PdfWriter

from muxi.runtime.services import observability
from muxi.runtime.services.multimodal.document_converter import (
    ANYDOC_EXTENSIONS,
    ATTACHMENT_CONVERTIBLE_EXTENSIONS,
    ENGINE_ANYDOC,
    ENGINE_MARKITDOWN,
    ENGINE_PDF_INSPECTOR,
    ConversionResult,
    QuarantineReason,
    _select_engine,
    _select_fallback_engine,
    convert_document,
    convert_document_async,
)

# ---------------------------------------------------------------------------
# Fixture builders (deterministic, generated in-test from real libraries)
# ---------------------------------------------------------------------------


def make_pdf_bytes(text: str = "Hello PDF World") -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(text=text)
    return bytes(pdf.output())


def make_docx_bytes(text: str = "Hello DOCX World") -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_bomb_docx_bytes(expanded_mb: int = 40) -> bytes:
    """A valid .docx whose document.xml holds enormous highly-compressible text."""
    base = make_docx_bytes("SEED")
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as zin:
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                data = zin.read(item)
                if item == "word/document.xml":
                    data = data.replace(b"SEED", b"A" * (expanded_mb * 1024 * 1024))
                zout.writestr(item, data)
    return buffer.getvalue()


def make_encrypted_pdf_bytes() -> bytes:
    reader = PdfReader(io.BytesIO(make_pdf_bytes("secret contents")))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("secret123")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


# A structurally-broken PDF (invalid xref table): pdf-inspector's strict parser
# rejects it, while MarkItDown's pdfminer backend recovers the text - the
# natural fixture for the pdf-inspector -> MarkItDown fallback path.
BROKEN_XREF_PDF = (
    b"%PDF-1.4\n"
    b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
    b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
    b"/Resources<</Font<</F1 5 0 R>>>>>>endobj\n"
    b"4 0 obj<</Length 44>>stream\n"
    b"BT /F1 24 Tf 72 700 Td (Hello Broken Xref) Tj ET\nendstream\nendobj\n"
    b"5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n"
    b"xref\n0 6\ntrailer<</Size 6/Root 1 0 R>>\nstartxref\n0\n%%EOF"
)


# ---------------------------------------------------------------------------
# Oversize input: rejected before any subprocess is spawned
# ---------------------------------------------------------------------------


def test_oversize_input_rejected_pre_spawn():
    result = convert_document(b"x" * 2048, "big.pdf", max_input_bytes=1024)
    assert not result.ok
    assert result.quarantine_reason == QuarantineReason.OVERSIZE
    # engine is only assigned when a worker is spawned; None proves the input
    # was refused before routing/spawning.
    assert result.engine is None
    assert "limit" in result.detail


# ---------------------------------------------------------------------------
# Well-formed documents: identical output to the previous in-process path
# ---------------------------------------------------------------------------


def _in_process_markitdown(content: bytes, suffix: str, tmp_path) -> str:
    path = tmp_path / f"reference{suffix}"
    path.write_bytes(content)
    return MarkItDown().convert(str(path)).text_content


def test_wellformed_docx_converts_via_anydoc():
    content = make_docx_bytes("Quarterly revenue grew 14% in Q3.")
    result = convert_document(content, "report.docx")
    assert result.ok, result.detail
    assert result.engine == ENGINE_ANYDOC
    assert not result.fallback_used
    assert "Quarterly revenue grew 14% in Q3." in result.text


def test_wellformed_html_converts_identically_to_in_process(tmp_path):
    content = b"<html><body><h1>Title</h1><p>hello sandboxed html</p></body></html>"
    result = convert_document(content, "page.html")
    assert result.ok, result.detail
    assert result.engine == ENGINE_MARKITDOWN
    assert result.text == _in_process_markitdown(content, ".html", tmp_path)


# ---------------------------------------------------------------------------
# PDF routing: pdf-inspector first, sandboxed MarkItDown as fallback
# ---------------------------------------------------------------------------


def test_pdf_routes_to_pdf_inspector():
    result = convert_document(make_pdf_bytes(), "hello.pdf", media_type="application/pdf")
    assert result.ok, result.detail
    assert result.engine == ENGINE_PDF_INSPECTOR
    assert not result.fallback_used
    assert "Hello PDF World" in result.text
    # Implementation marker: pdf-inspector emits structured markdown (heading
    # syntax); plain MarkItDown/pdfminer output for this file has no markdown.
    assert result.text.lstrip().startswith("#")


def test_pdf_media_type_routes_even_without_extension():
    result = convert_document(make_pdf_bytes(), "attachment", media_type="application/pdf")
    assert result.ok, result.detail
    assert result.engine == ENGINE_PDF_INSPECTOR


def test_pdf_inspector_failure_falls_back_to_markitdown():
    result = convert_document(BROKEN_XREF_PDF, "broken.pdf", media_type="application/pdf")
    assert result.ok, result.detail
    assert result.fallback_used
    assert result.engine == ENGINE_MARKITDOWN
    assert "Hello Broken Xref" in result.text
    # The pdf-inspector error is preserved for observability.
    assert "PDF parsing error" in result.detail


def test_garbage_pdf_quarantined_as_parser_error():
    result = convert_document(b"%PDF-1.7\n" + b"\x00\xff" * 200, "garbage.pdf")
    assert not result.ok
    assert result.quarantine_reason == QuarantineReason.PARSER_ERROR
    # Both engines were given a chance before quarantining.
    assert result.fallback_used


def test_encrypted_pdf_quarantined_as_encrypted():
    result = convert_document(make_encrypted_pdf_bytes(), "locked.pdf")
    assert not result.ok
    assert result.quarantine_reason == QuarantineReason.ENCRYPTED


# ---------------------------------------------------------------------------
# Routing table: anydoc is primary for its full claimed surface
# ---------------------------------------------------------------------------

# Every extension anydoc claims (minus .pdf, which keeps its dedicated route).
_ANYDOC_CLAIMED_EXTENSIONS = [
    ".doc",
    ".docx",
    ".docm",
    ".ppt",
    ".pps",
    ".pot",
    ".pptx",
    ".pptm",
    ".ppsx",
    ".ppsm",
    ".xls",
    ".xlsx",
    ".xlsm",
    ".xlsb",
    ".odt",
    ".ods",
    ".odp",
    ".rtf",
    ".epub",
    ".csv",
]


def test_anydoc_extension_set_matches_claimed_surface():
    assert ANYDOC_EXTENSIONS == frozenset(_ANYDOC_CLAIMED_EXTENSIONS)


@pytest.mark.parametrize("ext", _ANYDOC_CLAIMED_EXTENSIONS)
def test_every_claimed_extension_routes_to_anydoc(ext):
    assert _select_engine(f"file{ext}", None) == ENGINE_ANYDOC


@pytest.mark.parametrize(
    "media_type,expected",
    [
        ("application/pdf", ENGINE_PDF_INSPECTOR),
        ("application/msword", ENGINE_ANYDOC),
        (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ENGINE_ANYDOC,
        ),
        ("application/vnd.ms-word.document.macroEnabled.12", ENGINE_ANYDOC),
        ("application/vnd.ms-powerpoint", ENGINE_ANYDOC),
        (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            ENGINE_ANYDOC,
        ),
        ("application/vnd.openxmlformats-officedocument.presentationml.slideshow", ENGINE_ANYDOC),
        ("application/vnd.ms-excel", ENGINE_ANYDOC),
        ("application/vnd.ms-excel.sheet.binary.macroEnabled.12", ENGINE_ANYDOC),
        ("application/vnd.oasis.opendocument.text", ENGINE_ANYDOC),
        ("application/vnd.oasis.opendocument.spreadsheet", ENGINE_ANYDOC),
        ("application/vnd.oasis.opendocument.presentation", ENGINE_ANYDOC),
        ("application/rtf", ENGINE_ANYDOC),
        ("text/rtf", ENGINE_ANYDOC),
        ("application/epub+zip", ENGINE_ANYDOC),
        ("text/csv; charset=utf-8", ENGINE_ANYDOC),
        ("text/html", ENGINE_MARKITDOWN),
        ("application/octet-stream", ENGINE_MARKITDOWN),
        (None, ENGINE_MARKITDOWN),
    ],
)
def test_media_type_routing_without_extension(media_type, expected):
    # Extension-less filename: the media type alone must drive routing.
    assert _select_engine("attachment", media_type) == expected


def test_unclaimed_types_still_route_to_markitdown():
    assert _select_engine("page.html", None) == ENGINE_MARKITDOWN
    assert _select_engine("data.json", None) == ENGINE_MARKITDOWN
    assert _select_engine("mystery.bin", None) == ENGINE_MARKITDOWN
    assert _select_engine("doc.pdf", None) == ENGINE_PDF_INSPECTOR


@pytest.mark.parametrize(
    "filename,media_type,expected",
    [
        # pdf-inspector always falls back to MarkItDown.
        ("doc.pdf", "application/pdf", ENGINE_MARKITDOWN),
        # anydoc formats MarkItDown also supports keep a per-file fallback.
        ("a.docx", None, ENGINE_MARKITDOWN),
        ("a.pptx", None, ENGINE_MARKITDOWN),
        ("a.xlsx", None, ENGINE_MARKITDOWN),
        ("a.xls", None, ENGINE_MARKITDOWN),
        ("a.csv", None, ENGINE_MARKITDOWN),
        ("a.epub", None, ENGINE_MARKITDOWN),
        # Media type alone selects the fallback for extension-less files.
        ("attachment", "text/csv", ENGINE_MARKITDOWN),
        (
            "attachment",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ENGINE_MARKITDOWN,
        ),
        # anydoc-only formats have no fallback: failures quarantine honestly.
        ("a.doc", None, None),
        ("a.docm", None, None),
        ("a.ppt", None, None),
        ("a.pptm", None, None),
        ("a.xlsm", None, None),
        ("a.xlsb", None, None),
        ("a.odt", None, None),
        ("a.ods", None, None),
        ("a.odp", None, None),
        ("a.rtf", None, None),
        ("attachment", "application/msword", None),
        ("attachment", "application/vnd.oasis.opendocument.text", None),
    ],
)
def test_fallback_selection(filename, media_type, expected):
    engine = _select_engine(filename, media_type)
    assert _select_fallback_engine(engine, filename, media_type) == expected


def test_markitdown_route_has_no_fallback():
    assert _select_fallback_engine(ENGINE_MARKITDOWN, "page.html", "text/html") is None


def test_attachment_gate_covers_anydoc_surface():
    # Chat attachments must admit every format anydoc converts, plus the
    # pre-existing .pdf/.html routes.
    assert set(ATTACHMENT_CONVERTIBLE_EXTENSIONS) == {".pdf", ".html"} | ANYDOC_EXTENSIONS


# ---------------------------------------------------------------------------
# anydoc conversions: real fixtures for RTF / ODT / CSV / DOCX
# ---------------------------------------------------------------------------

RTF_FIXTURE = (
    rb"{\rtf1\ansi\deff0 {\fonttbl {\f0 Times New Roman;}}"
    rb"\f0\fs24 Hello {\b RTF} sandbox test.\par}"
)


def make_odt_bytes(text: str = "Hello ODT sandbox test.") -> bytes:
    """A minimal but valid ODT: a zip with the ODF mimetype and content.xml."""
    content_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        "<office:document-content"
        ' xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
        ' office:version="1.2">'
        f"<office:body><office:text><text:p>{text}</text:p></office:text></office:body>"
        "</office:document-content>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "mimetype",
            "application/vnd.oasis.opendocument.text",
            compress_type=zipfile.ZIP_STORED,
        )
        archive.writestr("content.xml", content_xml)
    return buffer.getvalue()


def test_rtf_converts_via_anydoc():
    result = convert_document(RTF_FIXTURE, "note.rtf")
    assert result.ok, result.detail
    assert result.engine == ENGINE_ANYDOC
    assert not result.fallback_used
    # anydoc renders the bold run as markdown emphasis.
    assert "Hello **RTF** sandbox test." in result.text


def test_odt_converts_via_anydoc():
    result = convert_document(make_odt_bytes(), "letter.odt")
    assert result.ok, result.detail
    assert result.engine == ENGINE_ANYDOC
    assert "Hello ODT sandbox test." in result.text


def test_csv_converts_via_anydoc_as_table():
    result = convert_document(b"name,qty\nwidget,3\ngadget,7\n", "data.csv", media_type="text/csv")
    assert result.ok, result.detail
    assert result.engine == ENGINE_ANYDOC
    # anydoc renders CSV as a markdown table.
    assert "| widget | 3 |" in result.text


def test_anydoc_failure_falls_back_to_markitdown_for_supported_format():
    # CSV bytes wearing an .xls extension: anydoc detects no OLE/ZIP signature,
    # trusts the extension, and its workbook parser rejects the file, while
    # MarkItDown's content sniffing recovers the table. This mirrors the
    # real-world servers-send-CSV-as-.xls case.
    result = convert_document(b"name,qty\nwidget,3\n", "export.xls")
    assert result.ok, result.detail
    assert result.fallback_used
    assert result.fallback_from == ENGINE_ANYDOC
    assert result.engine == ENGINE_MARKITDOWN
    assert "widget" in result.text
    # The anydoc error is preserved for observability.
    assert "ConvertError" in result.detail


def test_corrupt_anydoc_only_format_quarantines_without_fallback():
    # A structurally-broken ODT (zip local-file magic, no archive): odt is
    # anydoc-only, so the failure must quarantine honestly - no fallback.
    result = convert_document(b"PK\x03\x04 not really a zip archive", "broken.odt")
    assert not result.ok
    assert result.engine == ENGINE_ANYDOC
    assert result.quarantine_reason == QuarantineReason.PARSER_ERROR
    assert not result.fallback_used
    assert result.fallback_from is None


def test_fallback_event_payload_carries_primary_and_fallback_converters():
    capturing = _CapturingLogger()
    previous = observability.get_runtime_event_logger()
    observability.set_runtime_event_logger(capturing)
    observability.enable()
    try:
        convert_document(b"name,qty\nwidget,3\n", "export.xls")
        import time

        deadline = time.time() + 5
        fallback_events = []
        while time.time() < deadline:
            fallback_events = [
                e for e in capturing.snapshot() if e["event"] == "document.conversion.fallback"
            ]
            if fallback_events:
                break
            time.sleep(0.1)
    finally:
        observability.disable()
        observability.set_runtime_event_logger(previous)

    assert fallback_events, "no document.conversion.fallback event was emitted"
    data = fallback_events[0]["data"]
    assert data["primary"] == ENGINE_ANYDOC
    assert data["fallback"] == ENGINE_MARKITDOWN
    # Pre-anydoc field kept for existing consumers.
    assert data["fallback_from"] == ENGINE_ANYDOC


# ---------------------------------------------------------------------------
# Hostile fixtures: bombs and runaway parses are killed cleanly
# ---------------------------------------------------------------------------


@pytest.mark.timeout(180)
def test_decompression_bomb_docx_quarantined():
    """Tiny zip expanding to tens of MB of text must not reach the caller."""
    bomb = make_bomb_docx_bytes(expanded_mb=40)
    # The bomb passes the input-size gate easily (it is ~100KB on disk) ...
    assert len(bomb) < 1024 * 1024
    result = convert_document(bomb, "bomb.docx", max_output_bytes=1024 * 1024)
    # ... but the expanded output trips the sandbox: the explicit output cap
    # (oversize) or, on memory-limited platforms, the address-space limit.
    assert not result.ok
    assert result.quarantine_reason in (QuarantineReason.OVERSIZE, QuarantineReason.MEMORY)
    assert isinstance(result, ConversionResult)


@pytest.mark.timeout(120)
def test_runaway_parse_killed_at_wall_clock_timeout():
    """A parse that cannot finish in time is killed cleanly with TIMEOUT."""
    huge_html = b"<html><body>" + b"<div><p>word one two</p></div>" * 1_000_000 + b"</body></html>"
    result = convert_document(huge_html, "huge.html", timeout=1.5)
    assert not result.ok
    # Either kill mechanism is a valid clean kill: the parent's wall-clock
    # deadline ("exceeded ... wall-clock limit") or the worker's RLIMIT_CPU
    # backstop (SIGXCPU), whichever fires first. Both classify as TIMEOUT.
    assert result.quarantine_reason == QuarantineReason.TIMEOUT
    assert "wall-clock" in result.detail or "signal" in result.detail


# ---------------------------------------------------------------------------
# Service contract: never raises, async wrapper, quarantine event emission
# ---------------------------------------------------------------------------


def test_knowledge_loader_skips_oversize_file_before_reading(tmp_path):
    """FileKnowledge enforces max_file_size by stat BEFORE reading the bytes.

    The vector-ingestion path always had a pre-read size gate; the direct
    _load_file_content path (used by the reasoning tree builder) must apply
    the same guard instead of allocating the whole file first.
    """
    from muxi.runtime.formation.agents.knowledge.base import FileKnowledge

    big = tmp_path / "big.docx"
    big.write_bytes(b"x" * 4096)
    source = FileKnowledge(path=str(big), max_file_size=1024)
    content = source._load_file_content(str(big))
    assert content == "[Error loading file: big.docx]"


def test_async_wrapper_converts():
    result = asyncio.run(convert_document_async(make_pdf_bytes(), "hello.pdf"))
    assert result.ok
    assert result.engine == ENGINE_PDF_INSPECTOR


class _CapturingLogger:
    """Minimal event sink capturing what observe() emits (same shape as e2e 18c)."""

    def __init__(self):
        self.events = []
        self._lock = threading.Lock()

    def should_emit(self, event_type, level):
        return True

    def emit_event(self, event_type=None, level=None, data=None, description=None, **kwargs):
        with self._lock:
            self.events.append(
                {"event": getattr(event_type, "value", str(event_type)), "data": data}
            )
        return ""

    def snapshot(self):
        with self._lock:
            return list(self.events)


def test_quarantine_and_fallback_events_emitted():
    capturing = _CapturingLogger()
    previous = observability.get_runtime_event_logger()
    observability.set_runtime_event_logger(capturing)
    observability.enable()
    try:
        convert_document(b"x" * 2048, "big.pdf", max_input_bytes=1024)
        convert_document(BROKEN_XREF_PDF, "broken.pdf")
        # observe() emits on a short-lived background thread; give it time.
        import time

        deadline = time.time() + 5
        while time.time() < deadline:
            names = [e["event"] for e in capturing.snapshot()]
            if (
                "document.conversion.quarantined" in names
                and "document.conversion.fallback" in names
                and "document.conversion.completed" in names
            ):
                break
            time.sleep(0.1)
    finally:
        observability.disable()
        observability.set_runtime_event_logger(previous)

    events = capturing.snapshot()
    names = [e["event"] for e in events]
    assert "document.conversion.quarantined" in names
    assert "document.conversion.fallback" in names
    assert "document.conversion.completed" in names
    quarantined = next(e for e in events if e["event"] == "document.conversion.quarantined")
    assert quarantined["data"]["quarantine_reason"] == "oversize"
